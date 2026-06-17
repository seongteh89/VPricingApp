from decimal import Decimal

from openpyxl import load_workbook

from vpricing_app.models import ComparisonModel, Money, OriginalBQ, QuoteLine, SourceType, VendorQuote
from vpricing_app.services.filename_builder import build_output_filename
from vpricing_app.services.comparison_service import replace_vendor_quote
from vpricing_app.services.workbook_generator import export_comparison_workbook, load_comparison_metadata


def test_build_output_filename_includes_vendor_names():
    filename = build_output_filename(
        project_name="#Fire#AstraZeneca_F26B Fire Alarm System@Tuas South Ave",
        vendor_names=["Rensar", "E-Tech", "Rich"],
        revision=None,
    )

    assert filename.startswith("Quote Comparison - Fire AstraZeneca_F26B")
    assert "Rensar vs E-Tech vs Rich" in filename
    assert filename.endswith(".xlsx")


def test_build_output_filename_includes_revision_for_update():
    filename = build_output_filename("Sample Project", ["Rich"], revision="2026-06-16")

    assert filename == "Quote Comparison - Sample Project - Rich - Rev 2026-06-16.xlsx"


def test_export_workbook_contains_vendor_blocks_summary_and_metadata(tmp_path):
    original_line = QuoteLine(
        package="Package 1",
        section="A",
        item_no="1",
        description="Sub Alarm Panel",
        quantity=Decimal("5"),
        revised_quantity=Decimal("5"),
        unit="Nos",
    )
    vendor_line = original_line.model_copy(
        update={"labor_rate": Decimal("20"), "labor_total": Decimal("100"), "total": Decimal("100")}
    )
    original = OriginalBQ(project_name="Sample Project", source_filename="original.xlsx", packages={"Package 1": [original_line]})
    vendor = VendorQuote(
        vendor_name="Rich",
        source_filename="rich.pdf",
        source_type=SourceType.PDF,
        packages={"Package 1": [vendor_line]},
        package_totals={"Package 1": Money(total=Decimal("100"))},
    )
    model = ComparisonModel(project_name="Sample Project", original=original, vendors=[vendor])

    output_path = export_comparison_workbook(model, tmp_path)

    assert output_path.name == "Quote Comparison - Sample Project - Rich.xlsx"
    wb = load_workbook(output_path, data_only=True)
    assert "Summary" in wb.sheetnames
    assert "Package 1" in wb.sheetnames
    assert "Change Log" in wb.sheetnames
    assert "Import Warnings" in wb.sheetnames
    assert "System Metadata" in wb.sheetnames
    assert wb["System Metadata"].sheet_state == "hidden"
    summary = wb["Summary"]
    assert summary["B4"].value == "show Total Amount of each Vendor"
    assert summary["C5"].value == "Rich"
    assert summary["B6"].value == "package 1"
    assert summary["C6"].value == 100
    assert summary["B7"].value == "Total Package Amount"
    assert summary["C7"].value == 100
    package = wb["Package 1"]
    assert package.cell(package.max_row, 1).value == "GRAND TOTAL"
    assert package.cell(package.max_row, 10).value == 0
    assert package.cell(package.max_row, 11).value == 0
    assert package.cell(package.max_row, 12).value == 100
    assert package.cell(package.max_row, 12).fill.fgColor.rgb == "FFD9D9D9"
    metadata = load_comparison_metadata(output_path)
    assert metadata.project_name == "Sample Project"
    assert metadata.vendors[0].vendor_name == "Rich"


def test_export_workbook_splits_large_metadata_across_rows(tmp_path):
    lines = [
        QuoteLine(
            package="Package 1",
            section="A",
            item_no=str(index),
            description=f"Long description {index} " + ("x" * 120),
            quantity=Decimal("1"),
            revised_quantity=Decimal("1"),
            unit="Lot",
        )
        for index in range(500)
    ]
    original = OriginalBQ(project_name="Sample Project", source_filename="original.xlsx", packages={"Package 1": lines})
    vendor = VendorQuote(vendor_name="Rich", source_filename="rich.pdf", source_type=SourceType.PDF)
    model = ComparisonModel(project_name="Sample Project", original=original, vendors=[vendor])

    output_path = export_comparison_workbook(model, tmp_path)
    metadata = load_comparison_metadata(output_path)

    assert len(metadata.original.packages["Package 1"]) == 500


def test_replace_vendor_quote_preserves_other_vendors():
    original = OriginalBQ(project_name="Sample Project", source_filename="original.xlsx", packages={"Package 1": []})
    first = VendorQuote(vendor_name="Rich", source_filename="rich-r1.pdf", source_type=SourceType.PDF, revision="R1")
    second = VendorQuote(vendor_name="Rensar", source_filename="rensar.pdf", source_type=SourceType.EXCEL)
    revised = VendorQuote(vendor_name="Rich", source_filename="rich-r2.pdf", source_type=SourceType.PDF, revision="R2")
    model = ComparisonModel(project_name="Sample Project", original=original, vendors=[first, second])

    updated = replace_vendor_quote(model, revised)

    assert [vendor.vendor_name for vendor in updated.vendors] == ["Rich", "Rensar"]
    assert updated.vendors[0].source_filename == "rich-r2.pdf"
    assert updated.vendors[0].revision == "R2"
    assert updated.vendors[1].source_filename == "rensar.pdf"
