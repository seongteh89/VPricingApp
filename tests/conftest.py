from pathlib import Path

import pytest
from openpyxl import Workbook


@pytest.fixture
def sample_original_bq(tmp_path: Path) -> Path:
    path = tmp_path / "Sample_original.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "PACKAGE 1 with Meter Run"
    ws.append(["PROJECT NAME : #Fire#Sample Project"])
    ws.append([])
    ws.append(["Package 1 - Buildings"])
    ws.append(["NORMAL ZONE"])
    ws.append([None, None, "R0", "R1"])
    ws.append(
        [
            "Size:",
            "DESCRIPTION",
            "Quantity",
            "Revised Quantity",
            "Unit",
            "Material Rate",
            "Labor Rate",
            "Total Material Cost",
            "Total Labor Cost",
            "Total Cost",
            "Remarks",
        ]
    )
    ws.append(["A", "ADDRESSABLE FIRE ALARM SYSTEM", None, None, None, None, None, 0, 0, 0, None])
    ws.append([1, "Sub Alarm Panel", 5, 5, "Nos", None, None, 0, 0, 0, None])
    ws.append([2, "Manual Call Point", 10, 10, "Nos", None, None, 0, 0, 0, None])
    ws.append(["GRAND TOTAL", None, None, None, None, None, None, 0, 0, 0, None])

    scope = wb.create_sheet("Mechanical Scope")
    scope.append(["S/N", "Scope", "Acknowledgement"])
    scope.append([1, "Comply with safety rules", None])
    wb.save(path)
    return path


@pytest.fixture
def sample_vendor_excel(tmp_path: Path) -> Path:
    path = tmp_path / "Rensar quote.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "PACKAGE 1 with Meter Run"
    ws.append(["PROJECT NAME : #Fire#Sample Project"])
    ws.append([])
    ws.append(["Package 1 - Buildings"])
    ws.append(["NORMAL ZONE"])
    ws.append([None, None, "R0", "R1"])
    ws.append(
        [
            "Size:",
            "DESCRIPTION",
            "Quantity",
            "Revised Quantity",
            "Unit",
            "Material Rate",
            "Labor Rate",
            "Total Material Cost",
            "Total Labor Cost",
            "Total Cost",
            "Remarks",
        ]
    )
    ws.append(["A", "ADDRESSABLE FIRE ALARM SYSTEM", None, None, None, None, None, 0, 100, 100, None])
    ws.append([1, "Sub Alarm Panel", 5, 5, "Nos", None, 20, 0, 100, 100, None])
    ws.append([2, "Manual Call Point", 10, 12, "Nos", None, 10, 0, 120, 120, None])
    ws.append([3, "Vendor Extra Item", 1, 1, "Lot", 50, 50, 50, 50, 100, None])
    ws.append(["GRAND TOTAL", None, None, None, None, None, None, 50, 270, 320, None])
    wb.save(path)
    return path


@pytest.fixture
def sample_original_rfq(tmp_path: Path) -> Path:
    path = tmp_path / "SIN12 RFQ Original.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Request For Quotation"
    ws.append([])
    ws.append([])
    ws.append([])
    ws.append([None, None, "Request for Quotation"])
    ws.append([])
    ws.append([None, "Equipment     :", None, "Fire Alarm System "])
    for _ in range(6):
        ws.append([])
    ws.append([None, None, None, None, None, None, "Year 1", None, None, "Year 2", None, None, "Year 3", None, "Total Amount"])
    ws.append([None, None, None, None, None, None, "Jan - Dec 2027", None, None, "Jan - Dec 2028", None, None, "Jan - Dec 2029"])
    ws.append(
        [
            None,
            "No.",
            "Description",
            None,
            "Quantity",
            "Number of Servicing",
            "Unit Price",
            "Amount",
            "Number of Servicing",
            "Unit Price",
            "Amount",
            "Number of Servicing",
            "Unit Price",
            "Amount",
        ]
    )
    ws.append([])
    ws.append([None, None, "Preventive Maintenance"])
    ws.append([])
    ws.append([])
    ws.append([])
    ws.append([])
    ws.append([None, None, "PBB : Automatic Fire Sprinkler System", None, "1 Lot", 12])
    ws.append([None, None, "PBB : Fire Alarm System (Addressable)", None, 2, 12])
    ws.append([])
    ws.append([None, "Total", None, None, None, None, None, 0, None, None, 0, None, None, 0])
    wb.save(path)
    return path


@pytest.fixture
def sample_vendor_rfq(tmp_path: Path) -> Path:
    path = tmp_path / "SIN12 RFQ T-Tech.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Request For Quotation"
    ws.append([])
    ws.append([])
    ws.append([])
    ws.append([None, None, "Request for Quotation"])
    ws.append([])
    ws.append([None, "Equipment     :", None, "Fire Alarm System "])
    for _ in range(6):
        ws.append([])
    ws.append([None, None, None, None, None, None, "Year 1", None, None, "Year 2", None, None, "Year 3", None, "Total Amount"])
    ws.append([None, None, None, None, None, None, "Jan - Dec 2027", None, None, "Jan - Dec 2028", None, None, "Jan - Dec 2029"])
    ws.append(
        [
            None,
            "No.",
            "Description",
            None,
            "Quantity",
            "Number of Servicing",
            "Unit Price",
            "Amount",
            "Number of Servicing",
            "Unit Price",
            "Amount",
            "Number of Servicing",
            "Unit Price",
            "Amount",
        ]
    )
    ws.append([])
    ws.append([None, None, "Preventive Maintenance"])
    ws.append([])
    ws.append([])
    ws.append([])
    ws.append([])
    ws.append(
        [
            None,
            None,
            "PBB : Automatic Fire Sprinkler System",
            None,
            "1 Lot",
            12,
            600,
            7200,
            12,
            618,
            7416,
            12,
            636.54,
            7638.48,
        ]
    )
    ws.append(
        [
            None,
            None,
            "PBB : Fire Alarm System (Addressable)",
            None,
            2,
            12,
            400,
            4800,
            12,
            412,
            4944,
            12,
            424.36,
            5092.32,
        ]
    )
    ws.append([None, None, "Unpriced child location", None, None, None, None, 0, None, None, 0, None, None, 0])
    ws.append([])
    ws.append([None, "Total", None, None, None, None, None, 12000, None, None, 12360, None, None, 12730.8])
    wb.save(path)
    return path


@pytest.fixture
def sample_vendor_pdf(tmp_path: Path) -> Path:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import landscape, A4
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle

    path = tmp_path / "Rich quote.pdf"
    doc = SimpleDocTemplate(str(path), pagesize=landscape(A4))
    data = [
        ["PROJECT NAME : #Fire#Sample Project"],
        [
            "Package",
            "Section",
            "Item",
            "Description",
            "Qty",
            "Rev Qty",
            "Unit",
            "Mat Rate",
            "Lab Rate",
            "Mat Total",
            "Lab Total",
            "Total",
        ],
        ["PACKAGE 1 with Meter Run", "A", "1", "Sub Alarm Panel", "5", "5", "Nos", "", "15", "0", "75", "75"],
        ["PACKAGE 1 with Meter Run", "A", "2", "Manual Call Point", "10", "10", "Nos", "", "8", "0", "80", "80"],
        ["PACKAGE 1 with Meter Run", "B", "1", "Vendor Added Testing", "1", "1", "Lot", "25", "25", "25", "25", "50"],
        ["GRAND TOTAL", "", "", "", "", "", "", "", "", "25", "180", "205"],
    ]
    table = Table(data, repeatRows=1)
    table.setStyle(
        TableStyle(
            [
                ("GRID", (0, 0), (-1, -1), 0.25, colors.black),
                ("BACKGROUND", (0, 1), (-1, 1), colors.lightgrey),
            ]
        )
    )
    doc.build([table])
    return path


@pytest.fixture
def sample_vendor_pdf_bq_layout(tmp_path: Path) -> Path:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import landscape, A4
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle

    path = tmp_path / "Rich bq layout quote.pdf"
    doc = SimpleDocTemplate(str(path), pagesize=landscape(A4))
    data = [
        ["PROJECT NAME : #Fire#Sample Project", "", "", "", "", "", "", "", "", ""],
        ["Package 1 - Buildings", "", "Drug Product", "", "", "", "", "", "", ""],
        ["Size:", "DESCRIPTION", "Quantity", "Revised Quantity", "Unit", "Material Rate", "Labor Rate", "Total Material Cost", "Total Labor Cost", "Total Cost"],
        ["A", "ADDRESSABLE FIRE ALARM SYSTEM", "", "", "", "", "", "0", "75", "75"],
        ["1", "Sub Alarm Panel", "5", "5", "Nos", "", "15", "0", "75", "75"],
        ["GRAND TOTAL", "", "", "", "", "", "", "0", "75", "75"],
    ]
    table = Table(data, repeatRows=1)
    table.setStyle(
        TableStyle(
            [
                ("GRID", (0, 0), (-1, -1), 0.25, colors.black),
                ("BACKGROUND", (0, 2), (-1, 2), colors.lightgrey),
            ]
        )
    )
    doc.build([table])
    return path
