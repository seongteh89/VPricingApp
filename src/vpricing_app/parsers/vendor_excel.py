from pathlib import Path

from openpyxl import load_workbook

from vpricing_app.models import Money, QuoteLine, SourceType, VendorQuote, parse_money
from vpricing_app.parsers.original_bq import (
    _all_header_columns,
    _cell_value,
    _find_header_row,
    _first_header_column,
    _has_any_value,
    _header_columns,
    _is_package_sheet,
    _is_total_marker,
    _quantity_and_unit,
)


def _amount_value(amount: object, count: object, price: object):
    parsed_amount = parse_money(amount)
    if amount not in (None, ""):
        return parsed_amount
    return parse_money(count) * parse_money(price)


def parse_vendor_excel(path: str | Path, vendor_name: str) -> VendorQuote:
    wb = load_workbook(path, data_only=True)
    quote = VendorQuote(
        vendor_name=vendor_name,
        source_filename=Path(path).name,
        source_path=str(Path(path)),
        source_type=SourceType.EXCEL,
        is_official=False,
    )

    for ws in wb.worksheets:
        if not _is_package_sheet(ws.title):
            continue
        header_row = _find_header_row(ws)
        if header_row is None:
            continue
        headers = _header_columns(ws, header_row)
        item_col = _first_header_column(headers, "size", "no")
        description_col = _first_header_column(headers, "description")
        quantity_col = _first_header_column(headers, "quantity")
        revised_quantity_col = _first_header_column(headers, "revised quantity", "number of servicing")
        unit_col = _first_header_column(headers, "unit")
        material_rate_col = _first_header_column(headers, "material rate")
        labor_rate_col = _first_header_column(headers, "labor rate")
        material_total_col = _first_header_column(headers, "total material cost", "material total")
        labor_total_col = _first_header_column(headers, "total labor cost", "labor total")
        total_col = _first_header_column(headers, "total cost", "total")
        service_cols = _all_header_columns(headers, "number of servicing")
        unit_price_cols = _all_header_columns(headers, "unit price")
        amount_cols = _all_header_columns(headers, "amount")
        is_quotation_layout = bool(amount_cols and unit_price_cols and total_col is None)

        current_section: str | None = None
        current_section_description: str | None = None
        lines: list[QuoteLine] = []
        for row_number in range(header_row + 1, ws.max_row + 1):
            item_no = _cell_value(ws, row_number, item_col)
            description = _cell_value(ws, row_number, description_col)
            if _is_total_marker(item_no) or _is_total_marker(description):
                if is_quotation_layout:
                    totals = [parse_money(_cell_value(ws, row_number, column)) for column in amount_cols]
                    quote.package_totals[ws.title] = Money(
                        material=totals[0] if totals else parse_money(None),
                        labor=sum(totals[1:], parse_money(None)),
                        total=sum(totals, parse_money(None)),
                    )
                else:
                    quote.package_totals[ws.title] = Money(
                        material=parse_money(_cell_value(ws, row_number, material_total_col)),
                        labor=parse_money(_cell_value(ws, row_number, labor_total_col)),
                        total=parse_money(_cell_value(ws, row_number, total_col)),
                    )
                break
            if not description:
                continue
            description_text = str(description).strip()
            if isinstance(item_no, str) and len(item_no.strip()) == 1 and item_no.strip().isalpha():
                current_section = item_no.strip()
                current_section_description = description_text
                continue
            value_columns = (
                [quantity_col, revised_quantity_col, *service_cols, *unit_price_cols]
                if is_quotation_layout
                else [
                    quantity_col,
                    revised_quantity_col,
                    unit_col,
                    material_rate_col,
                    labor_rate_col,
                    material_total_col,
                    labor_total_col,
                    total_col,
                ]
            )
            if not _has_any_value(ws, row_number, value_columns):
                current_section_description = description_text
                continue
            quantity, unit = _quantity_and_unit(_cell_value(ws, row_number, quantity_col), _cell_value(ws, row_number, unit_col))
            if is_quotation_layout:
                amount_values = [
                    _amount_value(
                        _cell_value(ws, row_number, amount_col),
                        _cell_value(ws, row_number, service_cols[index] if index < len(service_cols) else None),
                        _cell_value(ws, row_number, unit_price_cols[index] if index < len(unit_price_cols) else None),
                    )
                    for index, amount_col in enumerate(amount_cols)
                ]
                material_rate = parse_money(_cell_value(ws, row_number, unit_price_cols[0] if unit_price_cols else None))
                labor_rate = parse_money(_cell_value(ws, row_number, unit_price_cols[1] if len(unit_price_cols) > 1 else None))
                material_total = amount_values[0] if amount_values else parse_money(None)
                labor_total = sum(amount_values[1:], parse_money(None))
                total = sum(amount_values, parse_money(None))
            else:
                material_rate = parse_money(_cell_value(ws, row_number, material_rate_col))
                labor_rate = parse_money(_cell_value(ws, row_number, labor_rate_col))
                material_total = parse_money(_cell_value(ws, row_number, material_total_col))
                labor_total = parse_money(_cell_value(ws, row_number, labor_total_col))
                total = parse_money(_cell_value(ws, row_number, total_col))
            lines.append(
                QuoteLine(
                    package=ws.title,
                    section=current_section,
                    section_description=current_section_description,
                    item_no=str(item_no).strip() if item_no is not None else None,
                    description=description_text,
                    quantity=quantity,
                    revised_quantity=parse_money(_cell_value(ws, row_number, revised_quantity_col)),
                    unit=unit,
                    material_rate=material_rate,
                    labor_rate=labor_rate,
                    material_total=material_total,
                    labor_total=labor_total,
                    total=total,
                    source_row=row_number,
                )
            )
        quote.packages[ws.title] = lines
        if ws.title not in quote.package_totals:
            quote.package_totals[ws.title] = Money(
                material=sum((line.material_total for line in lines), parse_money(None)),
                labor=sum((line.labor_total for line in lines), parse_money(None)),
                total=sum((line.total for line in lines), parse_money(None)),
            )

    return quote
