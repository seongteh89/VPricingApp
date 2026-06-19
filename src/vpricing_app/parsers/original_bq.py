import re
from pathlib import Path

from openpyxl import load_workbook

from vpricing_app.models import OriginalBQ, QuoteLine, parse_money


def _is_package_sheet(name: str) -> bool:
    normalized = name.strip().upper()
    return normalized.startswith("PACKAGE") or "QUOTATION" in normalized or "RFQ" in normalized


def _normalize_header(value: object) -> str:
    return " ".join(str(value or "").replace("\n", " ").strip().lower().split()).strip(": .")


def _header_columns(ws, row_number: int) -> dict[str, list[int]]:
    result: dict[str, list[int]] = {}
    for column in range(1, ws.max_column + 1):
        header = _normalize_header(ws.cell(row_number, column).value)
        if header:
            result.setdefault(header, []).append(column)
    return result


def _first_header_column(headers: dict[str, list[int]], *names: str) -> int | None:
    for name in names:
        columns = headers.get(name)
        if columns:
            return columns[0]
    return None


def _all_header_columns(headers: dict[str, list[int]], *names: str) -> list[int]:
    columns: list[int] = []
    for name in names:
        columns.extend(headers.get(name, []))
    return columns


def _find_header_row(ws) -> int | None:
    for row_number in range(1, ws.max_row + 1):
        headers = _header_columns(ws, row_number)
        if "description" in headers:
            return row_number
    return None


def _cell_value(ws, row_number: int, column: int | None):
    if column is None:
        return None
    return ws.cell(row_number, column).value


def _has_any_value(ws, row_number: int, columns: list[int | None]) -> bool:
    for column in columns:
        if column is None:
            continue
        value = ws.cell(row_number, column).value
        if value not in (None, ""):
            return True
    return False


def _is_total_marker(value: object) -> bool:
    text = str(value or "").strip().upper()
    return text in {"TOTAL", "GRAND TOTAL"}


def _quantity_and_unit(quantity_value: object, unit_value: object) -> tuple[object, str | None]:
    unit = str(unit_value or "").strip() or None
    if unit is None and isinstance(quantity_value, str):
        match = re.match(r"^[\d\s,.\-]+(.+)$", quantity_value.strip())
        if match:
            unit = match.group(1).strip() or None
    return parse_money(quantity_value), unit


def _project_name(value: object) -> str:
    text = str(value or "").strip()
    if "PROJECT NAME" in text.upper() and ":" in text:
        return text.split(":", 1)[1].strip()
    return text


def parse_original_bq(path: str | Path) -> OriginalBQ:
    wb = load_workbook(path, data_only=True)
    packages: dict[str, list[QuoteLine]] = {}
    project_name = Path(path).stem

    for ws in wb.worksheets:
        if not _is_package_sheet(ws.title):
            continue
        if ws["A1"].value:
            project_name = _project_name(ws["A1"].value)

        header_row = _find_header_row(ws)
        if header_row is None:
            continue
        headers = _header_columns(ws, header_row)
        item_col = _first_header_column(headers, "size", "no")
        description_col = _first_header_column(headers, "description")
        quantity_col = _first_header_column(headers, "quantity")
        revised_quantity_col = _first_header_column(headers, "revised quantity", "number of servicing")
        unit_col = _first_header_column(headers, "unit")
        amount_cols = _all_header_columns(headers, "amount", "total material cost", "total labor cost", "total cost")

        current_section: str | None = None
        current_section_description: str | None = None
        lines: list[QuoteLine] = []
        for row_number in range(header_row + 1, ws.max_row + 1):
            item_no = _cell_value(ws, row_number, item_col)
            description = _cell_value(ws, row_number, description_col)
            if not description:
                continue
            description_text = str(description).strip()
            if _is_total_marker(item_no) or _is_total_marker(description_text):
                break
            if isinstance(item_no, str) and len(item_no.strip()) == 1 and item_no.strip().isalpha():
                current_section = item_no.strip()
                current_section_description = description_text
                continue
            if not _has_any_value(ws, row_number, [quantity_col, revised_quantity_col, unit_col, *amount_cols]):
                current_section_description = description_text
                continue
            quantity, unit = _quantity_and_unit(_cell_value(ws, row_number, quantity_col), _cell_value(ws, row_number, unit_col))
            lines.append(
                QuoteLine(
                    package=ws.title,
                    section=current_section,
                    section_description=current_section_description,
                    item_no=str(item_no).strip() if item_no is not None else None,
                    description=description_text,
                    quantity=quantity,
                    revised_quantity=parse_money(_cell_value(ws, row_number, revised_quantity_col)),
                    unit=unit,
                    source_row=row_number,
                )
            )
        packages[ws.title] = lines

    return OriginalBQ(project_name=project_name, source_filename=Path(path).name, packages=packages)
