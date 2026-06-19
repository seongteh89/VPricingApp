import json
import re
from copy import copy
from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.formula.translate import Translator
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from vpricing_app.models import ComparisonModel, Money, QuoteLine, VendorQuote
from vpricing_app.parsers.original_bq import (
    _all_header_columns,
    _find_header_row,
    _first_header_column,
    _header_columns,
    _normalize_header,
)
from vpricing_app.services.filename_builder import build_output_filename
from vpricing_app.services.matcher import ComparisonRow, build_comparison_rows


HEADER_FILL = PatternFill("solid", fgColor="FFFFFFFF")
HEADER_FONT = Font(color="000000", bold=True)
ADDED_FILL = PatternFill("solid", fgColor="FFFFF2CC")
LOW_FILL = PatternFill("solid", fgColor="FFC6EFCE")
REVISED_QTY_FILL = PatternFill("solid", fgColor="FF7CC7E3")
MATERIAL_FILL = PatternFill("solid", fgColor="FF808080")
LABOR_FILL = PatternFill("solid", fgColor="FFFFFF00")
TOTAL_FILL = PatternFill("solid", fgColor="FFD9D9D9")
THIN_BORDER = Border(
    left=Side(style="thin", color="000000"),
    right=Side(style="thin", color="000000"),
    top=Side(style="thin", color="000000"),
    bottom=Side(style="thin", color="000000"),
)
MONEY_FORMAT = r'_-"$"* #,##0.00_-;\-"$"* #,##0.00_-;_-"$"* "-"??_-;_-@_-'
METADATA_MARKER = "comparison_model_json_v1"
METADATA_CHUNK_SIZE = 30000


def _as_float(value: Decimal | int | float | None) -> float:
    return float(value or 0)


def _summary_package_label(package_name: str) -> str:
    match = re.search(r"package\s*(\d+)", package_name, flags=re.IGNORECASE)
    return f"package {match.group(1)}" if match else package_name


def _package_total(vendor: VendorQuote, package_name: str) -> Money:
    if package_name in vendor.package_totals:
        return vendor.package_totals[package_name]
    lines = vendor.packages.get(package_name, [])
    return Money(
        material=sum((line.material_total for line in lines), Decimal("0")),
        labor=sum((line.labor_total for line in lines), Decimal("0")),
        total=sum((line.total for line in lines), Decimal("0")),
    )


def _write_headers(ws, vendors: list[str]) -> None:
    base_headers = ["Item", "Description", "Quantity", "Revised Quantity", "Unit"]
    vendor_headers = [
        "Revised Quantity",
        "Unit",
        "Material Rate",
        "Labor Rate",
        "Material Total",
        "Labor Total",
        "Total Cost",
        "Review Flag",
    ]
    headers = base_headers[:]
    for vendor in vendors:
        headers.extend([f"{vendor} {header}" for header in vendor_headers])
    ws.append(headers)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = THIN_BORDER


def _vendor_block_start(vendor_index: int) -> int:
    return 6 + vendor_index * 8


def _style_money_cell(cell, fill: PatternFill | None = None) -> None:
    cell.number_format = MONEY_FORMAT
    cell.border = THIN_BORDER
    if fill:
        cell.fill = fill


def _style_data_row(ws, row_number: int, vendor_count: int) -> None:
    for cell in ws[row_number]:
        cell.border = THIN_BORDER
    ws.cell(row_number, 3).fill = REVISED_QTY_FILL
    ws.cell(row_number, 4).fill = REVISED_QTY_FILL
    for vendor_index in range(vendor_count):
        start_col = _vendor_block_start(vendor_index)
        ws.cell(row_number, start_col).fill = REVISED_QTY_FILL
        ws.cell(row_number, start_col + 2).fill = MATERIAL_FILL
        ws.cell(row_number, start_col + 3).fill = LABOR_FILL
        ws.cell(row_number, start_col + 4).fill = MATERIAL_FILL
        for col in range(start_col + 2, start_col + 7):
            _style_money_cell(ws.cell(row_number, col))


def _reference_line(row: ComparisonRow) -> QuoteLine | None:
    if row.original_line:
        return row.original_line
    return next(iter(row.vendor_lines.values()), None)


def _section_key(row: ComparisonRow) -> tuple[str, str] | None:
    line = _reference_line(row)
    if not line or not line.section:
        return None
    return line.section, line.section_description or ""


def _write_section_subtotal_row(ws, section_key: tuple[str, str], vendor_count: int) -> int:
    section, description = section_key
    values = [section, description, "", "", ""]
    for _ in range(vendor_count):
        values.extend(["", "", "", "", "", "", "", ""])
    ws.append(values)
    row_number = ws.max_row
    for cell in ws[row_number]:
        cell.font = Font(bold=True)
        cell.fill = TOTAL_FILL
        cell.border = THIN_BORDER
    return row_number


def _write_section_subtotal_formulas(
    ws,
    subtotal_row: int | None,
    first_item_row: int | None,
    last_item_row: int | None,
    vendor_count: int,
) -> None:
    if subtotal_row is None or first_item_row is None or last_item_row is None:
        return
    for vendor_index in range(vendor_count):
        start_col = _vendor_block_start(vendor_index)
        for col in (start_col + 4, start_col + 5, start_col + 6):
            column_letter = get_column_letter(col)
            cell = ws.cell(subtotal_row, col)
            cell.value = f"=SUM({column_letter}{first_item_row}:{column_letter}{last_item_row})"
            cell.font = Font(bold=True)
            _style_money_cell(cell, TOTAL_FILL)


def _write_summary(summary, model: ComparisonModel) -> None:
    summary.append(["Vendor", "Grand Total"])
    for cell in summary[1]:
        cell.font = HEADER_FONT
        cell.border = THIN_BORDER
    for vendor in model.vendors:
        summary.append([vendor.vendor_name, _as_float(vendor.grand_total)])
        summary.cell(summary.max_row, 2).number_format = MONEY_FORMAT

    block_start = len(model.vendors) + 3
    summary.cell(block_start, 2).value = "show Total Amount of each Vendor"
    header_row = block_start + 1
    for index, vendor in enumerate(model.vendors, start=3):
        cell = summary.cell(header_row, index)
        cell.value = vendor.vendor_name
        cell.font = HEADER_FONT
        cell.border = THIN_BORDER

    package_start = header_row + 1
    for row_offset, package_name in enumerate(model.original.packages):
        row_number = package_start + row_offset
        summary.cell(row_number, 2).value = _summary_package_label(package_name)
        for col_offset, vendor in enumerate(model.vendors, start=3):
            cell = summary.cell(row_number, col_offset)
            cell.value = _as_float(_package_total(vendor, package_name).total)
            cell.number_format = MONEY_FORMAT

    total_row = package_start + len(model.original.packages)
    summary.cell(total_row, 2).value = "Total Package Amount"
    summary.cell(total_row, 2).font = Font(bold=True)
    for col_offset, vendor in enumerate(model.vendors, start=3):
        total = sum((_package_total(vendor, package_name).total for package_name in model.original.packages), Decimal("0"))
        cell = summary.cell(total_row, col_offset)
        cell.value = _as_float(total)
        cell.font = Font(bold=True)
        cell.number_format = MONEY_FORMAT
        cell.fill = TOTAL_FILL


def _write_package_total_row(ws, package_name: str, vendors: list[VendorQuote]) -> None:
    values = ["GRAND TOTAL", "", "", "", ""]
    for vendor in vendors:
        total = _package_total(vendor, package_name)
        values.extend(["", "", "", "", _as_float(total.material), _as_float(total.labor), _as_float(total.total), ""])
    ws.append(values)
    row_number = ws.max_row
    for cell in ws[row_number]:
        cell.font = Font(bold=True)
        cell.fill = TOTAL_FILL
        cell.border = THIN_BORDER
    for vendor_index in range(len(vendors)):
        start_col = _vendor_block_start(vendor_index)
        for col in (start_col + 4, start_col + 5, start_col + 6):
            _style_money_cell(ws.cell(row_number, col), TOTAL_FILL)


def _autosize(ws) -> None:
    for column in ws.columns:
        width = max(len(str(cell.value or "")) for cell in column)
        ws.column_dimensions[get_column_letter(column[0].column)].width = min(max(width + 2, 10), 45)


def _package_number(name: str) -> str | None:
    match = re.search(r"package\s*(\d+)", name, flags=re.IGNORECASE)
    return match.group(1) if match else None


def _worksheet_for_package(wb, package_name: str):
    if package_name in wb.sheetnames:
        return wb[package_name]
    package_number = _package_number(package_name)
    if package_number is not None:
        for ws in wb.worksheets:
            if _package_number(ws.title) == package_number:
                return ws
    return None


def _worksheet_for_vendor_package(wb, package_name: str, vendor_name: str):
    exact_titles = [f"{package_name} - {vendor_name}"]
    package_number = _package_number(package_name)
    if package_number is not None:
        exact_titles.append(f"Package {package_number} - {vendor_name}")
    for title in exact_titles:
        if title in wb.sheetnames:
            return wb[title]

    normalized_vendor = vendor_name.casefold()
    for ws in wb.worksheets:
        if normalized_vendor not in ws.title.casefold():
            continue
        if package_number is None or _package_number(ws.title) == package_number:
            return ws
    return None


def _first_total_amount_column(ws, header_row: int, start_column: int) -> int | None:
    for row_number in range(1, header_row + 1):
        for column in range(start_column, ws.max_column + 1):
            if _normalize_header(ws.cell(row_number, column).value) == "total amount":
                return column
    return None


def _pricing_columns(ws) -> list[int]:
    header_row = _find_header_row(ws)
    if header_row is None:
        return []
    headers = _header_columns(ws, header_row)
    start_candidates = [
        _first_header_column(headers, "revised quantity"),
        _first_header_column(headers, "number of servicing"),
    ]
    start_columns = [column for column in start_candidates if column is not None]
    if not start_columns:
        return []
    start_column = min(start_columns)

    total_amount_column = _first_total_amount_column(ws, header_row, start_column)
    if total_amount_column is not None:
        return list(range(start_column, total_amount_column + 1))

    pricing_candidates = [
        _first_header_column(headers, "revised quantity"),
        _first_header_column(headers, "unit"),
        _first_header_column(headers, "material rate"),
        _first_header_column(headers, "labor rate"),
        _first_header_column(headers, "total material cost", "material total"),
        _first_header_column(headers, "total labor cost", "labor total"),
        _first_header_column(headers, "total cost", "total"),
        *_all_header_columns(headers, "number of servicing"),
        *_all_header_columns(headers, "unit price"),
        *_all_header_columns(headers, "amount"),
    ]
    columns = [column for column in pricing_candidates if column is not None]
    if not columns:
        return []
    columns = [column for column in columns if column >= start_column]
    return list(range(start_column, max(columns) + 1))


def _columns_in_block(headers: dict[str, list[int]], block_columns: set[int], *names: str) -> list[int]:
    columns: list[int] = []
    for name in names:
        columns.extend(column for column in headers.get(name, []) if column in block_columns)
    return columns


def _first_column_in_block(headers: dict[str, list[int]], block_columns: set[int], *names: str) -> int | None:
    columns = _columns_in_block(headers, block_columns, *names)
    return columns[0] if columns else None


def _pricing_column_map(ws) -> dict[str, int | list[int] | None]:
    header_row = _find_header_row(ws)
    if header_row is None:
        return {}
    headers = _header_columns(ws, header_row)
    block_columns = set(_pricing_columns(ws))
    return {
        "revised_quantity": _first_column_in_block(headers, block_columns, "revised quantity", "number of servicing"),
        "unit": _first_column_in_block(headers, block_columns, "unit"),
        "material_rate": _first_column_in_block(headers, block_columns, "material rate"),
        "labor_rate": _first_column_in_block(headers, block_columns, "labor rate"),
        "material_total": _first_column_in_block(headers, block_columns, "total material cost", "material total"),
        "labor_total": _first_column_in_block(headers, block_columns, "total labor cost", "labor total"),
        "total": _first_column_in_block(headers, block_columns, "total cost", "total"),
        "service_cols": _columns_in_block(headers, block_columns, "number of servicing"),
        "unit_price_cols": _columns_in_block(headers, block_columns, "unit price"),
        "amount_cols": _columns_in_block(headers, block_columns, "amount"),
    }


def _shift_pricing_column(value: int | list[int] | None, offset: int) -> int | list[int] | None:
    if isinstance(value, list):
        return [column + offset for column in value]
    if isinstance(value, int):
        return value + offset
    return value


def _shift_pricing_map(
    pricing_map: dict[str, int | list[int] | None],
    offset: int,
) -> dict[str, int | list[int] | None]:
    return {key: _shift_pricing_column(value, offset) for key, value in pricing_map.items()}


def _write_cell_if_column(ws, row_number: int, column: int | None, value) -> None:
    if column is not None and value is not None:
        ws.cell(row_number, column).value = _as_float(value) if isinstance(value, Decimal) else value


def _write_parsed_pricing(target_ws, target_row: int, line: QuoteLine, pricing_map: dict[str, int | list[int] | None]) -> None:
    service_cols = pricing_map.get("service_cols") or []
    unit_price_cols = pricing_map.get("unit_price_cols") or []
    amount_cols = pricing_map.get("amount_cols") or []
    if isinstance(service_cols, list) and isinstance(unit_price_cols, list) and isinstance(amount_cols, list) and amount_cols:
        _write_cell_if_column(target_ws, target_row, service_cols[0] if service_cols else None, line.revised_quantity)
        _write_cell_if_column(target_ws, target_row, unit_price_cols[0] if unit_price_cols else None, line.material_rate)
        _write_cell_if_column(target_ws, target_row, amount_cols[0], line.material_total)
        if len(unit_price_cols) > 1:
            _write_cell_if_column(target_ws, target_row, unit_price_cols[1], line.labor_rate)
        if len(amount_cols) > 1:
            _write_cell_if_column(target_ws, target_row, amount_cols[1], line.labor_total)
        if len(amount_cols) > 2:
            _write_cell_if_column(target_ws, target_row, amount_cols[-1], line.total)
        return

    _write_cell_if_column(target_ws, target_row, pricing_map.get("revised_quantity"), line.revised_quantity)
    _write_cell_if_column(target_ws, target_row, pricing_map.get("unit"), line.unit)
    _write_cell_if_column(target_ws, target_row, pricing_map.get("material_rate"), line.material_rate)
    _write_cell_if_column(target_ws, target_row, pricing_map.get("labor_rate"), line.labor_rate)
    _write_cell_if_column(target_ws, target_row, pricing_map.get("material_total"), line.material_total)
    _write_cell_if_column(target_ws, target_row, pricing_map.get("labor_total"), line.labor_total)
    _write_cell_if_column(target_ws, target_row, pricing_map.get("total"), line.total)


def _translated_cell_value(value, source_coordinate: str, target_coordinate: str):
    if isinstance(value, str) and value.startswith("="):
        try:
            return Translator(value, origin=source_coordinate).translate_formula(target_coordinate)
        except Exception:
            return value
    return value


def _copy_template_cell(source, target) -> None:
    target.value = _translated_cell_value(source.value, source.coordinate, target.coordinate)
    if source.has_style:
        target._style = copy(source._style)
    if source.hyperlink:
        target._hyperlink = copy(source.hyperlink)
    if source.comment:
        target.comment = copy(source.comment)


def _copy_pricing_cell(source_ws, source_row: int, target_ws, target_row: int, source_column: int, target_column: int) -> None:
    source = source_ws.cell(source_row, source_column)
    target = target_ws.cell(target_row, target_column)
    target.value = _translated_cell_value(source.value, source.coordinate, target.coordinate)


def _vendor_pricing_columns(pricing_columns: list[int], vendor_index: int) -> list[int]:
    offset = vendor_index * len(pricing_columns)
    return [column + offset for column in pricing_columns]


def _merged_range_overlaps(ws, min_col: int, min_row: int, max_col: int, max_row: int) -> bool:
    for merged_range in ws.merged_cells.ranges:
        existing_min_col, existing_min_row, existing_max_col, existing_max_row = merged_range.bounds
        columns_overlap = min_col <= existing_max_col and max_col >= existing_min_col
        rows_overlap = min_row <= existing_max_row and max_row >= existing_min_row
        if columns_overlap and rows_overlap:
            return True
    return False


def _merge_range_if_possible(ws, min_col: int, min_row: int, max_col: int, max_row: int) -> None:
    if min_col == max_col and min_row == max_row:
        return
    if _merged_range_overlaps(ws, min_col, min_row, max_col, max_row):
        return
    ws.merge_cells(
        start_row=min_row,
        start_column=min_col,
        end_row=max_row,
        end_column=max_col,
    )


def _copy_pricing_block(ws, source_columns: list[int], target_columns: list[int]) -> None:
    if not source_columns or not target_columns:
        return
    offset = target_columns[0] - source_columns[0]
    for source_column, target_column in zip(source_columns, target_columns, strict=True):
        source_letter = get_column_letter(source_column)
        target_letter = get_column_letter(target_column)
        source_dimension = ws.column_dimensions[source_letter]
        target_dimension = ws.column_dimensions[target_letter]
        target_dimension.width = source_dimension.width
        target_dimension.hidden = source_dimension.hidden
        target_dimension.outlineLevel = source_dimension.outlineLevel
        for row_number in range(1, ws.max_row + 1):
            _copy_template_cell(ws.cell(row_number, source_column), ws.cell(row_number, target_column))

    first_source_column = min(source_columns)
    last_source_column = max(source_columns)
    for merged_range in list(ws.merged_cells.ranges):
        min_col, min_row, max_col, max_row = merged_range.bounds
        if min_col < first_source_column or max_col > last_source_column:
            continue
        _merge_range_if_possible(ws, min_col + offset, min_row, max_col + offset, max_row)


def _ensure_vendor_pricing_blocks(ws, pricing_columns: list[int], vendor_count: int) -> None:
    if vendor_count <= 1 or not pricing_columns:
        return
    block_width = len(pricing_columns)
    ws.insert_cols(max(pricing_columns) + 1, amount=(vendor_count - 1) * block_width)
    for vendor_index in range(1, vendor_count):
        _copy_pricing_block(ws, pricing_columns, _vendor_pricing_columns(pricing_columns, vendor_index))


def _vendor_label_row(ws, target_columns: list[int]) -> int | None:
    header_row = _find_header_row(ws)
    if header_row is None or not target_columns:
        return None
    first_column = min(target_columns)
    last_column = max(target_columns)
    for row_number in range(header_row - 1, 0, -1):
        first_value = ws.cell(row_number, first_column).value
        other_values = [ws.cell(row_number, column).value for column in range(first_column + 1, last_column + 1)]
        if first_value not in (None, "") and all(value in (None, "") for value in other_values):
            return row_number
    for row_number in range(header_row - 1, 0, -1):
        values = [ws.cell(row_number, column).value for column in range(first_column, last_column + 1)]
        if all(value in (None, "") for value in values):
            return row_number
    return max(header_row - 1, 1)


def _write_vendor_block_label(ws, vendor_name: str, target_columns: list[int]) -> None:
    label_row = _vendor_label_row(ws, target_columns)
    if label_row is None:
        return
    first_column = min(target_columns)
    last_column = max(target_columns)
    cell = ws.cell(label_row, first_column)
    cell.value = vendor_name
    cell.font = Font(bold=True)
    cell.alignment = Alignment(horizontal="center", vertical="center")
    _merge_range_if_possible(ws, first_column, label_row, last_column, label_row)


def _copy_vendor_pricing(
    source_ws,
    target_ws,
    rows: list[ComparisonRow],
    vendor: VendorQuote,
    source_columns: list[int],
    target_columns: list[int],
) -> None:
    offset = target_columns[0] - source_columns[0] if source_columns and target_columns else 0
    pricing_map = _shift_pricing_map(_pricing_column_map(target_ws), offset)
    for row in rows:
        vendor_line = row.vendor_lines.get(vendor.vendor_name)
        if vendor_line is None:
            continue
        target_row = row.original_line.source_row if row.original_line and row.original_line.source_row else vendor_line.source_row
        if target_row is None:
            continue
        if source_ws is not None and source_columns and target_columns:
            if vendor_line.source_row is None:
                _write_parsed_pricing(target_ws, target_row, vendor_line, pricing_map)
                continue
            for source_column, target_column in zip(source_columns, target_columns, strict=True):
                _copy_pricing_cell(source_ws, vendor_line.source_row, target_ws, target_row, source_column, target_column)
        else:
            _write_parsed_pricing(target_ws, target_row, vendor_line, pricing_map)


def _replace_sheet(wb, title: str):
    if title in wb.sheetnames:
        del wb[title]
    return wb.create_sheet(title)


def _write_change_log_sheet(wb, model: ComparisonModel) -> None:
    change_log = _replace_sheet(wb, "Change Log")
    change_log.append(["Level", "Code", "Vendor", "Package", "Row", "Message"])
    for warning in model.warnings:
        change_log.append(
            [warning.level.value, warning.code, warning.vendor_name, warning.package, warning.row_label, warning.message]
        )
    _autosize(change_log)


def _write_import_warnings_sheet(wb, model: ComparisonModel) -> None:
    warnings = _replace_sheet(wb, "Import Warnings")
    warnings.append(["Level", "Code", "Vendor", "Package", "Message"])
    for warning in model.warnings:
        warnings.append([warning.level.value, warning.code, warning.vendor_name, warning.package, warning.message])
    _autosize(warnings)


def _write_metadata_sheet(wb, model: ComparisonModel) -> None:
    metadata = _replace_sheet(wb, "System Metadata")
    metadata.sheet_state = "hidden"
    metadata["A1"] = METADATA_MARKER
    raw_metadata = model.model_dump_json()
    for index in range(0, len(raw_metadata), METADATA_CHUNK_SIZE):
        row = index // METADATA_CHUNK_SIZE + 2
        metadata.cell(row=row, column=1).value = raw_metadata[index : index + METADATA_CHUNK_SIZE]


def _write_support_sheets(wb, model: ComparisonModel) -> None:
    _write_change_log_sheet(wb, model)
    _write_import_warnings_sheet(wb, model)
    _write_metadata_sheet(wb, model)


def _export_template_preserved_workbook(
    model: ComparisonModel,
    output_path: Path,
    original_template_path: str | Path,
) -> Path:
    wb = load_workbook(original_template_path)
    rows_by_package = build_comparison_rows(model.original, model.vendors)
    source_workbooks = {
        vendor.vendor_name: load_workbook(vendor.source_path, data_only=False)
        for vendor in model.vendors
        if vendor.source_path and Path(vendor.source_path).suffix.lower() in {".xlsx", ".xlsm"}
    }

    for package_name, rows in rows_by_package.items():
        target_ws = _worksheet_for_package(wb, package_name)
        if target_ws is None:
            continue
        pricing_columns = _pricing_columns(target_ws)
        _ensure_vendor_pricing_blocks(target_ws, pricing_columns, len(model.vendors))
        for vendor_index, vendor in enumerate(model.vendors):
            target_columns = _vendor_pricing_columns(pricing_columns, vendor_index)
            _write_vendor_block_label(target_ws, vendor.vendor_name, target_columns)
            source_wb = source_workbooks.get(vendor.vendor_name)
            source_ws = _worksheet_for_package(source_wb, package_name) if source_wb is not None else None
            _copy_vendor_pricing(source_ws, target_ws, rows, vendor, pricing_columns, target_columns)

    if "Summary" in wb.sheetnames:
        del wb["Summary"]
    summary = wb.create_sheet("Summary", 0)
    _write_summary(summary, model)
    _autosize(summary)
    _write_support_sheets(wb, model)
    wb.save(output_path)
    return output_path


def export_updated_template_workbook(
    model: ComparisonModel,
    existing_comparison_path: str | Path,
    revised_vendor: VendorQuote,
    output_dir: str | Path,
    revision: str | None = None,
) -> Path:
    output_dir = Path(output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)
    filename = build_output_filename(model.project_name, [vendor.vendor_name for vendor in model.vendors], revision)
    output_path = output_dir / filename

    wb = load_workbook(existing_comparison_path)
    rows_by_package = build_comparison_rows(model.original, model.vendors)
    source_wb = load_workbook(revised_vendor.source_path, data_only=False) if revised_vendor.source_path else None
    updated_any_sheet = False

    for package_name, rows in rows_by_package.items():
        target_ws = _worksheet_for_package(wb, package_name)
        if target_ws is None:
            target_ws = _worksheet_for_vendor_package(wb, package_name, revised_vendor.vendor_name)
        if target_ws is None:
            continue
        source_ws = _worksheet_for_package(source_wb, package_name) if source_wb is not None else None
        pricing_columns = _pricing_columns(target_ws)
        vendor_index = next(
            (index for index, vendor in enumerate(model.vendors) if vendor.vendor_name == revised_vendor.vendor_name),
            0,
        )
        target_columns = _vendor_pricing_columns(pricing_columns, vendor_index)
        _write_vendor_block_label(target_ws, revised_vendor.vendor_name, target_columns)
        _copy_vendor_pricing(source_ws, target_ws, rows, revised_vendor, pricing_columns, target_columns)
        updated_any_sheet = True

    if not updated_any_sheet:
        return export_comparison_workbook(model, output_dir, revision=revision)

    if "Summary" in wb.sheetnames:
        del wb["Summary"]
    summary = wb.create_sheet("Summary", 0)
    _write_summary(summary, model)
    _autosize(summary)
    _write_support_sheets(wb, model)
    wb.save(output_path)
    return output_path


def export_comparison_workbook(
    model: ComparisonModel,
    output_dir: str | Path,
    revision: str | None = None,
    original_template_path: str | Path | None = None,
) -> Path:
    output_dir = Path(output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)
    filename = build_output_filename(model.project_name, [vendor.vendor_name for vendor in model.vendors], revision)
    output_path = output_dir / filename

    if original_template_path is not None:
        return _export_template_preserved_workbook(model, output_path, original_template_path)

    wb = Workbook()
    summary = wb.active
    summary.title = "Summary"
    _write_summary(summary, model)
    _autosize(summary)

    rows_by_package = build_comparison_rows(model.original, model.vendors)
    for package_name, rows in rows_by_package.items():
        ws = wb.create_sheet(package_name[:31])
        _write_headers(ws, [vendor.vendor_name for vendor in model.vendors])
        current_section_key: tuple[str, str] | None = None
        subtotal_row: int | None = None
        first_item_row: int | None = None
        last_item_row: int | None = None
        for row in rows:
            section_key = _section_key(row)
            if section_key != current_section_key:
                _write_section_subtotal_formulas(
                    ws, subtotal_row, first_item_row, last_item_row, len(model.vendors)
                )
                current_section_key = section_key
                subtotal_row = None
                first_item_row = None
                last_item_row = None
                if section_key is not None:
                    subtotal_row = _write_section_subtotal_row(ws, section_key, len(model.vendors))

            original = row.original_line
            first_vendor_line = next(iter(row.vendor_lines.values()), None)
            values = [
                original.item_no if original else "",
                original.description if original else first_vendor_line.description,
                float(original.quantity) if original and original.quantity is not None else "",
                float(original.revised_quantity) if original and original.revised_quantity is not None else "",
                original.unit if original else "",
            ]
            vendor_totals = []
            for vendor in model.vendors:
                vendor_line = row.vendor_lines.get(vendor.vendor_name)
                if vendor_line is None:
                    values.extend(["N/A", "N/A", "", "", "", "", "", "Missing"])
                    vendor_totals.append(None)
                    continue
                values.extend(
                    [
                        float(vendor_line.revised_quantity) if vendor_line.revised_quantity is not None else "",
                        vendor_line.unit or "",
                        float(vendor_line.material_rate) if vendor_line.material_rate is not None else "",
                        float(vendor_line.labor_rate) if vendor_line.labor_rate is not None else "",
                        float(vendor_line.material_total),
                        float(vendor_line.labor_total),
                        float(vendor_line.total),
                        "Vendor Added Item - Review Required" if vendor_line.is_vendor_added else "",
                    ]
                )
                vendor_totals.append(vendor_line.total)
            ws.append(values)
            _style_data_row(ws, ws.max_row, len(model.vendors))
            if original is None:
                for cell in ws[ws.max_row]:
                    cell.fill = ADDED_FILL
            numeric_totals = [total for total in vendor_totals if total is not None]
            if numeric_totals:
                lowest = min(numeric_totals)
                for index, total in enumerate(vendor_totals):
                    if total == lowest:
                        total_col = 5 + index * 8 + 7
                        ws.cell(ws.max_row, total_col).fill = LOW_FILL
            if current_section_key is not None:
                first_item_row = first_item_row or ws.max_row
                last_item_row = ws.max_row
        _write_section_subtotal_formulas(ws, subtotal_row, first_item_row, last_item_row, len(model.vendors))
        _write_package_total_row(ws, package_name, model.vendors)
        _autosize(ws)

    _write_support_sheets(wb, model)

    wb.save(output_path)
    return output_path


def load_comparison_metadata(path: str | Path) -> ComparisonModel:
    wb = load_workbook(path, data_only=True)
    sheet = wb["System Metadata"]
    if sheet["A1"].value == METADATA_MARKER:
        chunks = []
        row = 2
        while sheet.cell(row=row, column=1).value:
            chunks.append(sheet.cell(row=row, column=1).value)
            row += 1
        raw = "".join(chunks)
    else:
        raw = sheet["A1"].value
    return ComparisonModel.model_validate(json.loads(raw))
